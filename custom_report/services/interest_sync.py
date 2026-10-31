import os
import time
from datetime import datetime
import openpyxl
from tqdm import tqdm
import frappe
from frappe.utils import now_datetime
from custom_report.api.calculator import get_account_details, get_branch_code


def parse_date(str_val):
	if not str_val:
		return None
	if isinstance(str_val, datetime):
		return str_val
	s = str(str_val).strip()
	if "/" in s:
		p = s.split("/")
		if len(p) == 3:
			try:
				return datetime(int(p[2]), int(p[1]), int(p[0]))
			except Exception:
				pass
	if "-" in s:
		p = s.split("-")
		if len(p) == 3:
			try:
				if len(p[0]) == 4:
					return datetime(int(p[0]), int(p[1]), int(p[2]))
				return datetime(int(p[2]), int(p[1]), int(p[0]))
			except Exception:
				pass
	return None


def to_iso_date(dt):
	if not dt:
		return None
	return dt.strftime("%Y-%m-%d")


def calculate_maturity_date(opn_date_str, period_mths):
	start = parse_date(opn_date_str)
	try:
		mths = int(period_mths or 0)
	except Exception:
		mths = 0
	if not start or mths <= 0:
		return None
	target_month = start.month + mths
	target_year = start.year + (target_month - 1) // 12
	target_month = ((target_month - 1) % 12) + 1
	target_day = start.day

	if target_month == 12:
		next_month = datetime(target_year + 1, 1, 1)
	else:
		next_month = datetime(target_year, target_month + 1, 1)
	from dateutil.relativedelta import relativedelta
	last_day = (next_month - relativedelta(days=1)).day
	if target_day > last_day:
		target_day = last_day
	return datetime(target_year, target_month, target_day)


def get_canonical_scheme_code(c):
	c = str(c or "").strip().upper()
	if c in ["0010", "SMBG"]:
		return "2005"
	if c == "0040":
		return "2015"
	return c


SCHEME_RULES = {
	"2005": {"name": "SMBG (66M)", "tenure_m": 66, "base_rate": 8.0, "slabs": [[36, 66, 6.0]]},
	"2010": {"name": "RD 12M", "tenure_m": 12, "base_rate": 8.0, "slabs": [[6, 9, 4.0], [9, 12, 6.0]]},
	"2011": {"name": "RD 24M", "tenure_m": 24, "base_rate": 8.0, "slabs": [[12, 18, 4.0], [18, 24, 6.0]]},
	"2012": {"name": "RD 36M", "tenure_m": 36, "base_rate": 8.0, "slabs": [[18, 27, 4.0], [27, 36, 6.0]]},
	"2013": {"name": "RD 48M", "tenure_m": 48, "base_rate": 8.0, "slabs": [[24, 36, 4.0], [36, 48, 6.0]]},
	"2014": {"name": "RD 60M", "tenure_m": 60, "base_rate": 8.0, "slabs": [[30, 45, 4.0], [45, 60, 6.0]]},
	"2015": {"name": "RD 84M", "tenure_m": 84, "base_rate": 8.0, "slabs": [[42, 63, 4.0], [63, 84, 6.0]]},
	"2016": {"name": "RD 120M", "tenure_m": 120, "base_rate": 8.0, "slabs": [[60, 90, 4.0], [90, 120, 6.0]]}
}


def get_applicable_rules(sc, m, is_premature):
	sc = get_canonical_scheme_code(sc)
	r = SCHEME_RULES.get(sc)
	if not r:
		return 0.0
	if not is_premature:
		return r["base_rate"]
	ar = 0.0
	for s in r["slabs"]:
		if s[0] <= m < s[1]:
			ar = s[2]
			break
	if ar == 0.0 and m >= r["tenure_m"] and r["slabs"]:
		ar = r["slabs"][-1][2]
	return ar


def get_service_charge_percent(schm_code, months_held, is_premature=True):
	sc = get_canonical_scheme_code(schm_code)
	if sc == "2005":
		return 18.0 if is_premature else 0.0
	if sc == "2010":
		if months_held < 3:
			return 6.0
		if months_held < 6:
			return 4.0
		return 0.0
	return 0.0


def run_calculation(details, imported_transactions):
	opn_dt_obj = parse_date(details.get("acct_opn_date"))
	period_mths = details.get("deposit_period_mths") or 0
	mat_dt_obj = calculate_maturity_date(details.get("acct_opn_date"), period_mths)
	if not mat_dt_obj or not opn_dt_obj:
		return None

	today = datetime.now()
	exit_dt_obj = mat_dt_obj if today > mat_dt_obj else today
	is_premature = exit_dt_obj < mat_dt_obj

	db_txns = details.get("transactions") or []
	valid_csv = [t for t in imported_transactions if parse_date(t.get("date"))]
	csv_split_dt = parse_date("06/12/2024")
	valid_db = [t for t in db_txns if not (parse_date(t.get("date")) and csv_split_dt and parse_date(t.get("date")) <= csv_split_dt)]

	combined = list(valid_csv)
	for dt in valid_db:
		d_dt = parse_date(dt.get("date"))
		d_str = d_dt.strftime("%d/%m/%Y") if d_dt else ""
		d_amt = float(dt.get("amount") or 0)
		is_dup = any(
			(parse_date(c.get("date")).strftime("%d/%m/%Y") if parse_date(c.get("date")) else "") == d_str
			and abs(float(c.get("amount") or 0) - d_amt) < 0.01
			for c in valid_csv
		)
		if not is_dup:
			combined.append(dt)

	combined.sort(key=lambda x: parse_date(x.get("date")) or datetime.min)

	effective_txns = []
	for t in combined:
		if not t.get("is_exceptional_reversal"):
			amt = float(t.get("amount") or 0)
			t_type = str(t.get("type") or "").strip().upper()
			if amt < 0 or t_type == "D":
				break
		effective_txns.append(t)

	def get_cycle_offset(d):
		offset = (d.year - opn_dt_obj.year) * 12 + (d.month - opn_dt_obj.month)
		if offset > 0 and d.day <= opn_dt_obj.day:
			offset -= 1
		return max(offset, 0)

	raw_credits_by_month = {}
	actual_credits_total = 0.0
	max_txn_month = 0

	for t in effective_txns:
		d_val = parse_date(t.get("date"))
		if not d_val:
			continue
		amt = float(t.get("amount") or 0)
		t_type = str(t.get("type") or "").strip().upper()

		if t_type == "C":
			if d_val < mat_dt_obj and d_val <= exit_dt_obj:
				m_off = get_cycle_offset(d_val)
				if not t.get("is_exceptional_reversal"):
					actual_credits_total += amt
					raw_credits_by_month[m_off] = raw_credits_by_month.get(m_off, 0.0) + amt
					if amt > 0 and m_off > max_txn_month:
						max_txn_month = m_off
		elif t_type == "D" and t.get("is_exceptional_reversal"):
			if d_val <= exit_dt_obj:
				m_off = get_cycle_offset(d_val)
				debit_amt = abs(amt)
				reverses = not t.get("is_partial_withdrawal") and any(
					ot.get("type") == "C" and ot.get("is_exceptional_reversal")
					and abs(float(ot.get("amount") or 0) - debit_amt) < 0.01
					and parse_date(ot.get("date")) < d_val
					for ot in effective_txns
				)
				if not reverses:
					actual_credits_total -= debit_amt
					raw_credits_by_month[m_off] = raw_credits_by_month.get(m_off, 0.0) - debit_amt

	sett_m_offset = get_cycle_offset(exit_dt_obj)
	tenure_months = int(period_mths or 0)
	planned_installment = float(details.get("planned_installment") or details.get("deposit_amount") or 0)
	cycle_count = max(tenure_months, sett_m_offset + 1, max_txn_month + 1, 1)

	allocated = {}
	if planned_installment <= 0:
		for m in range(cycle_count):
			allocated[m] = raw_credits_by_month.get(m, 0.0)
	else:
		carry = 0.0
		for m in range(cycle_count):
			carry += raw_credits_by_month.get(m, 0.0)
			if carry + 0.0001 >= planned_installment:
				allocated[m] = planned_installment
				carry -= planned_installment
			else:
				allocated[m] = carry
				carry = 0.0

	m_rel = (exit_dt_obj.year - opn_dt_obj.year) * 12 + (exit_dt_obj.month - opn_dt_obj.month)
	if exit_dt_obj.day < opn_dt_obj.day:
		m_rel -= 1
	months_held_total = m_rel
	effective_interest_months = sett_m_offset + 1

	app_rate = get_applicable_rules(details.get("schm_code"), months_held_total, is_premature)

	principal_sum = 0.0
	compounded_int = 0.0
	accrued_curr_qtr = 0.0
	total_interest_accrued = 0.0

	max_allocated_month = 0
	for k, v in allocated.items():
		if v > 0 and k > max_allocated_month:
			max_allocated_month = k

	display_limit = (effective_interest_months - 1) if is_premature else max(max_allocated_month, tenure_months - 1)

	for m in range(display_limit + 1):
		inst = allocated.get(m, 0.0)
		if m < effective_interest_months:
			principal_sum += inst

		m_int = 0.0
		should_accrue = (m < effective_interest_months) and (m < tenure_months)
		interest_bearing_bal = (principal_sum if is_premature else (principal_sum + compounded_int)) > 0

		if should_accrue and interest_bearing_bal and (not is_premature or m < sett_m_offset):
			base_bal = principal_sum if is_premature else (principal_sum + compounded_int)
			m_int = base_bal * (app_rate / 1200.0)
			total_interest_accrued += m_int
			accrued_curr_qtr += m_int

		if not is_premature and (m + 1) % 3 == 0:
			compounded_int += accrued_curr_qtr
			accrued_curr_qtr = 0.0

	final_principal = round(actual_credits_total, 2)
	final_interest = round(total_interest_accrued)
	sc_percent = get_service_charge_percent(details.get("schm_code"), months_held_total, is_premature)
	penalty = round(((actual_credits_total * sc_percent) / 100.0) * 1.18)
	passbook = 0
	if get_canonical_scheme_code(details.get("schm_code")) == "2005" and months_held_total < 36:
		passbook = 60

	final_net = round(final_principal + final_interest - penalty - passbook)

	return {
		"opening_date": to_iso_date(opn_dt_obj),
		"closing_date": to_iso_date(mat_dt_obj),
		"interest": final_interest,
		"principal": final_principal,
		"finacle_maturity_amount": float(details.get("maturity_amount") or 0),
		"final_net_settlement": final_net
	}


def resolve_file_path(file_path=None):
	if file_path and os.path.exists(file_path):
		return file_path

	attached_file = frappe.db.get_single_value("Interest Settings", "interest_log_file")
	if not attached_file:
		# Fallback to local default path if available
		default_path = "/home/talib/Downloads/TD_Migrated.xlsx"
		if os.path.exists(default_path):
			return default_path
		raise ValueError("No file found! Please upload file in Interest Settings or provide file_path.")

	if attached_file.startswith("/"):
		full_path = frappe.get_site_path("public" if attached_file.startswith("/files/") else "", attached_file.lstrip("/"))
		if os.path.exists(full_path):
			return full_path

	file_doc = frappe.get_doc("File", {"file_url": attached_file})
	full_path = file_doc.get_full_path()
	if os.path.exists(full_path):
		return full_path

	raise FileNotFoundError(f"File not found on system: {attached_file}")


def sync_interest_logs(file_path=None, batch_size=50, limit=None):
	"""
	Execute with:
	bench --site sahayog.com execute custom_report.services.interest_sync.sync_interest_logs
	"""
	actual_file = resolve_file_path(file_path)
	print(f"\n[INFO] Starting Interest Sync using file: {actual_file}")

	wb = openpyxl.load_workbook(actual_file, read_only=True)
	sheet = wb.active

	# Fetch existing accounts to avoid redundant processing
	existing = set(frappe.get_all("Interest Log", fields=["account_number"], as_list=True))
	processed_accounts = {str(r[0]).strip() for r in existing if r and r[0]}
	print(f"[INFO] Existing records already in Interest Log: {len(processed_accounts)}")

	total_processed = 0
	success_count = 0
	failed_count = 0
	skipped_count = 0

	rows = sheet.iter_rows(values_only=True)
	headers = next(rows, None)
	if not headers:
		print("[ERROR] Excel file is empty.")
		return

	# Header map
	header_map = {str(h).strip().lower(): idx for idx, h in enumerate(headers) if h is not None}
	acc_idx = header_map.get("account_number", 0)

	start_time = time.time()
	total_rows = int(limit) if limit else ((sheet.max_row - 1) if sheet.max_row else None)
	pbar = tqdm(total=total_rows, desc="Calculating Interest", unit="rec", dynamic_ncols=True)

	try:
		for row_idx, row in enumerate(rows, start=1):
			if limit and total_processed >= int(limit):
				break

			raw_acc = row[acc_idx] if len(row) > acc_idx else None
			if not raw_acc:
				continue

			acc = str(raw_acc).split(".")[0].strip()
			if not acc:
				continue

			if acc in processed_accounts:
				skipped_count += 1
				pbar.update(1)
				pbar.set_postfix({"ok": success_count, "fail": failed_count, "skip": skipped_count})
				continue

			try:
				prefix = acc[:4]
				scheme_code = acc[4:8]
				netwin_ac = acc[8:]

				details = get_account_details(foracid=acc)
				if not details or not details.get("success"):
					err_msg = details.get("error") if details else "Failed to fetch Finacle details"
					frappe.get_doc({
						"doctype": "Interest Log",
						"account_number": acc,
						"status": "Failed",
						"error_message": str(err_msg)[:500]
					}).insert(ignore_permissions=True)
					failed_count += 1
				else:
					imported_txns = []
					if scheme_code in ["0010", "1540", "0040"]:
						b_res = get_branch_code(branch_prefix=prefix)
						branch_code = b_res.get("branch_code") if b_res.get("success") else ""
						gl_code = "B101" if scheme_code == "0010" else ("RD" if scheme_code == "1540" else ("B104" if scheme_code == "0040" else ""))
						if branch_code and gl_code and netwin_ac:
							try:
								opn_val = (details.get("acct_opn_date") or "01/01/1990").split("/")
								start_iso = f"{opn_val[2]}-{opn_val[1]}-{opn_val[0]}" if len(opn_val) == 3 else "1990-01-01"
								frappe.call("custom_report.api.netwin_statement.download_netwin_statement",
											branch_code=branch_code, ac_code=gl_code, ac_no=netwin_ac,
											start_date=start_iso, end_date="2099-12-31", export_format="csv")
								csv_content = frappe.local.response.get("filecontent", "")
								if csv_content:
									lines = [l.strip() for l in csv_content.splitlines() if l.strip()]
									hdr_idx = -1
									for li, l in enumerate(lines):
										if l.lower().startswith("date"):
											hdr_idx = li
											break
									if hdr_idx != -1:
										h_line = [h.strip().lower() for h in lines[hdr_idx].split(",")]
										d_idx = next((i for i, h in enumerate(h_line) if "date" in h), -1)
										cr_idx = next((i for i, h in enumerate(h_line) if "credit" in h or h == "cr"), -1)
										dr_idx = next((i for i, h in enumerate(h_line) if "debit" in h or h == "dr"), -1)
										desc_idx = next((i for i, h in enumerate(h_line) if "desc" in h or "particular" in h), -1)
										for line in lines[hdr_idx + 1:]:
											cols = [c.strip() for c in line.split(",")]
											if len(cols) > max(d_idx, cr_idx, dr_idx):
												cr_val = float(cols[cr_idx].replace(",", "") or 0) if cr_idx != -1 and cols[cr_idx] else 0
												dr_val = float(cols[dr_idx].replace(",", "") or 0) if dr_idx != -1 and cols[dr_idx] else 0
												p_desc = cols[desc_idx] if desc_idx != -1 else "Netwin Import"
												if cr_val > 0:
													imported_txns.append({"date": cols[d_idx], "amount": cr_val, "type": "C", "particular": p_desc, "imported_from_csv": True, "is_exceptional_reversal": False})
												elif dr_val > 0:
													imported_txns.append({"date": cols[d_idx], "amount": dr_val, "type": "D", "particular": p_desc, "imported_from_csv": True, "is_exceptional_reversal": True})
							except Exception:
								pass

					calc = run_calculation(details, imported_txns)
					if calc:
						frappe.get_doc({
							"doctype": "Interest Log",
							"account_number": acc,
							"status": "Success",
							"opening_date": calc.get("opening_date"),
							"closing_date": calc.get("closing_date"),
							"principle_amount": calc.get("principal", 0.0),
							"ineterest_calculated": calc.get("interest", 0.0),
							"maturity_amount": calc.get("finacle_maturity_amount", 0.0),
							"final_net_settlement": calc.get("final_net_settlement", 0.0)
						}).insert(ignore_permissions=True)
						success_count += 1
					else:
						frappe.get_doc({
							"doctype": "Interest Log",
							"account_number": acc,
							"status": "Failed",
							"error_message": "Calculation logic returned empty result"
						}).insert(ignore_permissions=True)
						failed_count += 1

			except Exception as e:
				frappe.get_doc({
					"doctype": "Interest Log",
					"account_number": acc,
					"status": "Failed",
					"error_message": str(e)[:500]
				}).insert(ignore_permissions=True)
				failed_count += 1

			processed_accounts.add(acc)
			total_processed += 1
			pbar.update(1)
			pbar.set_postfix({"ok": success_count, "fail": failed_count, "skip": skipped_count})

			# Commit in batches
			if total_processed % batch_size == 0:
				frappe.db.commit()

	except (KeyboardInterrupt, SystemExit):
		print("\n[STOP] Process stopped gracefully by user. Saving completed batch...")
	finally:
		frappe.db.commit()
		pbar.close()
		wb.close()
		print(f"\n[DONE] Summary: Processed: {total_processed} | Success: {success_count} | Failed: {failed_count} | Skipped: {skipped_count}")
